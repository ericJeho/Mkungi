"""Rebuild a page layout as an editable Excel workbook.

One sheet per page. Free text goes above/below tables in reading order;
tables keep their borders, merged cells and relative column widths, and
numeric text becomes real numbers so totals can be recalculated.
"""

from __future__ import annotations

import io
import re

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from app.models import DocumentLayout, Page, Table, TextLine

_THIN = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_NUMBER = re.compile(r"^\(?-?[A-Z]{0,3}\s?[\d,]+(\.\d+)?\)?%?$")


def parse_number(text: str) -> float | int | None:
    """'1,250.00' -> 1250.0, '(300)' -> -300, 'MK 5,000' -> 5000. Else None."""
    t = text.strip()
    if not t or not _NUMBER.match(t) or not any(ch.isdigit() for ch in t):
        return None
    negative = t.startswith("(") and t.endswith(")")
    pct = t.endswith("%")
    t = re.sub(r"[A-Z(),%\s]", "", t)
    try:
        value: float | int = float(t) if "." in t else int(t)
    except ValueError:
        return None
    if negative:
        value = -value
    if pct:
        value = value / 100
    return value


def _set(cell: Cell, text: str, bold: bool = False) -> None:
    number = parse_number(text)
    if number is not None:
        cell.value = number
        cell.number_format = "0.00%" if text.strip().endswith("%") else (
            "#,##0.00" if isinstance(number, float) else "#,##0"
        )
    else:
        cell.value = text
    cell.font = Font(bold=bold)
    cell.alignment = Alignment(wrap_text=True, vertical="top")


def _table_bold(page: Page, table: Table) -> set[tuple[int, int]]:
    bold = set()
    for c in table.cells:
        lines = [ln for ln in page.lines if c.bbox.contains_point(ln.bbox.cx, ln.bbox.cy)]
        if lines and all(ln.bold for ln in lines):
            bold.add((c.row, c.col))
    return bold


def _write_page(ws, page: Page) -> None:
    free = page.free_lines()
    items: list[tuple[int, object]] = [(ln.bbox.y0, ln) for ln in free] + [(t.bbox.y0, t) for t in page.tables]
    items.sort(key=lambda it: it[0])

    # Column grid: the widest table defines it; text uses column A onwards.
    widest = max(page.tables, key=lambda t: t.n_cols, default=None)
    if widest:
        total = sum(widest.col_widths) or 1
        for i, w in enumerate(widest.col_widths):
            ws.column_dimensions[get_column_letter(i + 1)].width = max(8, round(100 * w / total))
    else:
        ws.column_dimensions["A"].width = 80

    row = 1
    last_y: int | None = None
    pending: list[TextLine] = []

    def flush() -> None:
        nonlocal row
        for col, ln in enumerate(sorted(pending, key=lambda ln: ln.bbox.x0), start=1):
            _set(ws.cell(row=row, column=col), ln.text, ln.bold)
        if pending:
            row += 1
        pending.clear()

    for _, item in items:
        if isinstance(item, TextLine):
            if last_y is not None and pending and abs(item.bbox.y0 - last_y) > 0.5 * item.bbox.height:
                flush()
            pending.append(item)
            last_y = item.bbox.y0
            continue
        flush()
        table: Table = item  # type: ignore[assignment]
        bold = _table_bold(page, table)
        start = row
        for c in table.cells:
            cell = ws.cell(row=start + c.row, column=c.col + 1)
            _set(cell, c.text, (c.row, c.col) in bold)
            for k in range(c.col_span):
                ws.cell(row=start + c.row, column=c.col + 1 + k).border = _BORDER
            if c.col_span > 1:
                ws.merge_cells(
                    start_row=start + c.row,
                    start_column=c.col + 1,
                    end_row=start + c.row,
                    end_column=c.col + c.col_span,
                )
        row = start + table.n_rows + 1
        last_y = None
    flush()


def write_xlsx(layout: DocumentLayout) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for i, page in enumerate(layout.pages, start=1):
        _write_page(wb.create_sheet(f"Page {i}"), page)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
