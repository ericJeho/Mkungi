"""Rebuild a page layout as an editable Word document.

The output is flow layout (paragraphs, tables, inline pictures) rather than
absolutely positioned text boxes, so it stays easy to edit. Position is kept
through indents, tab stops, alignment and spacing derived from the scan.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from typing import Union

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor

from app.models import DocumentLayout, ImageRegion, Page, Table, TextLine

A4_SHORT_IN = 8.27
A4_LONG_IN = 11.69
MARGIN_IN = 0.5
EMU_PER_IN = 914400
MAX_GAP_PT = 600  # Word allows up to 1584pt of paragraph spacing


@dataclass
class _Row:
    """Text lines that share a baseline band, written as one paragraph."""

    lines: list[TextLine]

    @property
    def y0(self) -> int:
        return min(ln.bbox.y0 for ln in self.lines)

    @property
    def y1(self) -> int:
        return max(ln.bbox.y1 for ln in self.lines)


Element = Union[_Row, Table, ImageRegion]


def _group_rows(lines: list[TextLine]) -> list[_Row]:
    rows: list[_Row] = []
    for ln in lines:
        for row in rows:
            overlap = min(row.y1, ln.bbox.y1) - max(row.y0, ln.bbox.y0)
            if overlap > 0.5 * min(ln.bbox.height, row.y1 - row.y0):
                row.lines.append(ln)
                break
        else:
            rows.append(_Row([ln]))
    for row in rows:
        row.lines.sort(key=lambda ln: ln.bbox.x0)
    return rows


def _top(el: Element) -> int:
    return el.y0 if isinstance(el, _Row) else el.bbox.y0


def _bottom(el: Element) -> int:
    return el.y1 if isinstance(el, _Row) else el.bbox.y1


def _set_cell_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tbl_pr.append(borders)


class _PageWriter:
    def __init__(self, doc, page: Page, first: bool) -> None:
        self.doc = doc
        self.page = page
        landscape = page.width > page.height
        width_in = A4_LONG_IN if landscape else A4_SHORT_IN
        self.in_per_px = width_in / page.width

        section = doc.sections[0] if first else doc.add_section()
        section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        section.page_width = Inches(width_in)
        section.page_height = Inches(A4_SHORT_IN if landscape else A4_LONG_IN)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, side, Inches(MARGIN_IN))
        self.content_in = width_in - 2 * MARGIN_IN

    def x_in(self, px: float) -> float:
        """Page-pixel x to inches from the left margin."""
        return min(max(0.0, px * self.in_per_px - MARGIN_IN), self.content_in)

    def font_pt(self, line: TextLine) -> float:
        # An OCR line box is roughly 1.25x the font size.
        pt = line.bbox.height * self.in_per_px * 72 / 1.25
        return max(6.0, min(48.0, round(pt * 2) / 2))

    def write(self) -> None:
        elements: list[Element] = [*_group_rows(self.page.free_lines()), *self.page.tables, *self.page.images]
        elements.sort(key=_top)
        prev_bottom = int(MARGIN_IN / self.in_per_px)
        for el in elements:
            gap_pt = max(0.0, (_top(el) - prev_bottom) * self.in_per_px * 72)
            if isinstance(el, _Row):
                self._row(el, gap_pt)
            elif isinstance(el, Table):
                self._spacer(gap_pt)
                self._table(el)
            else:
                self._image(el, gap_pt)
            prev_bottom = max(prev_bottom, _bottom(el))

    def _spacer(self, gap_pt: float) -> None:
        if gap_pt > 4:
            p = self.doc.add_paragraph()
            fmt = p.paragraph_format
            fmt.space_before = Pt(0)
            fmt.space_after = Pt(0)
            fmt.line_spacing = Pt(min(gap_pt, MAX_GAP_PT))

    def _run(self, p, line: TextLine) -> None:
        run = p.add_run(line.text)
        run.font.size = Pt(self.font_pt(line))
        run.bold = line.bold
        if line.color != (0, 0, 0):
            run.font.color.rgb = RGBColor(*line.color)

    def _row(self, row: _Row, gap_pt: float) -> None:
        p = self.doc.add_paragraph()
        fmt = p.paragraph_format
        fmt.space_before = Pt(min(gap_pt, MAX_GAP_PT))
        fmt.space_after = Pt(0)
        first = row.lines[0]
        page_cx = self.page.width / 2

        if len(row.lines) == 1:
            # Centred lines (titles, letterheads) and right-aligned lines (dates, refs).
            if abs(first.bbox.cx - page_cx) < 0.04 * self.page.width and first.bbox.x0 > 0.2 * self.page.width:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif first.bbox.x1 > 0.85 * self.page.width and first.bbox.x0 > 0.5 * self.page.width:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            else:
                fmt.left_indent = Inches(self.x_in(first.bbox.x0))
            self._run(p, first)
            return

        # Several lines side by side: use tab stops at their x positions.
        fmt.left_indent = Inches(self.x_in(first.bbox.x0))
        self._run(p, first)
        for line in row.lines[1:]:
            fmt.tab_stops.add_tab_stop(Inches(self.x_in(line.bbox.x0)), WD_TAB_ALIGNMENT.LEFT)
            p.add_run("\t")
            self._run(p, line)

    def _table(self, table: Table) -> None:
        wt = self.doc.add_table(rows=table.n_rows, cols=table.n_cols)
        wt.alignment = WD_TABLE_ALIGNMENT.LEFT
        wt.autofit = False
        _set_cell_borders(wt)
        widths = [Emu(int(w * self.in_per_px * EMU_PER_IN)) for w in table.col_widths]
        for r, row in enumerate(wt.rows):
            row.height = Emu(int(table.row_heights[r] * self.in_per_px * EMU_PER_IN))
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            for i, cell in enumerate(row.cells):
                cell.width = widths[i]
        size = Pt(max(8.0, min(12.0, self._median_font(table))))
        for c in table.cells:
            cell = wt.cell(c.row, c.col)
            if c.col_span > 1:
                cell = cell.merge(wt.cell(c.row, c.col + c.col_span - 1))
            cell.text = ""
            para = cell.paragraphs[0]
            for i, text in enumerate(c.text.split("\n") if c.text else []):
                if i:
                    para.add_run().add_break()
                run = para.add_run(text)
                run.font.size = size
                run.bold = self._cell_bold(c)

    def _median_font(self, table: Table) -> float:
        sizes = sorted(
            self.font_pt(ln) for ln in self.page.lines if table.bbox.contains_point(ln.bbox.cx, ln.bbox.cy)
        )
        return sizes[len(sizes) // 2] if sizes else 10.0

    def _cell_bold(self, cell) -> bool:
        lines = [ln for ln in self.page.lines if cell.bbox.contains_point(ln.bbox.cx, ln.bbox.cy)]
        return bool(lines) and all(ln.bold for ln in lines)

    def _image(self, region: ImageRegion, gap_pt: float) -> None:
        p = self.doc.add_paragraph()
        fmt = p.paragraph_format
        fmt.space_before = Pt(min(gap_pt, MAX_GAP_PT))
        fmt.space_after = Pt(0)
        fmt.left_indent = Inches(self.x_in(region.bbox.x0))
        width = min(region.bbox.width * self.in_per_px, self.content_in)
        p.add_run().add_picture(io.BytesIO(region.png), width=Inches(width))


def write_docx(layout: DocumentLayout) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.paragraph_format.space_after = Pt(0)
    for i, page in enumerate(layout.pages):
        _PageWriter(doc, page, first=i == 0).write()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
