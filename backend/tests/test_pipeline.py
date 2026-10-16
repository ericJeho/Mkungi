import io
import time

import cv2
import numpy as np
import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.models import BBox, ImageKind, TextLine
from app.pipeline.classify import classify, recommended_format
from app.pipeline.layout import detect_tables
from app.pipeline.run import convert
from app.pipeline.writers.xlsx_writer import parse_number

W, H = 1240, 1754  # A4 at 150 dpi

# Table: header row spans all 3 columns, then 3 rows x 3 columns.
TX0, TX1 = 100, 1140
ROWS = [600, 660, 720, 780, 840]
COLS = [100, 600, 850, 1140]


def _invoice_image() -> np.ndarray:
    img = np.full((H, W, 3), 255, np.uint8)
    # Logo: solid blue block top-left. Stamp: red ring bottom-right.
    cv2.rectangle(img, (100, 80), (260, 200), (200, 80, 20), -1)
    cv2.circle(img, (1000, 1400), 90, (30, 30, 220), 8)
    for y in ROWS:
        cv2.line(img, (TX0, y), (TX1, y), (0, 0, 0), 2)
    for x in COLS:
        top = ROWS[0] if x in (TX0, TX1) else ROWS[1]  # no inner rules in the header row
        cv2.line(img, (x, top), (x, ROWS[-1]), (0, 0, 0), 2)
    for text, (x, y) in _TEXT.items():
        cv2.putText(img, text, (x, y), cv2.FONT_HERSHEY_SIMPLEX, 0.9, (0, 0, 0), 2)
    return img


_TEXT = {
    "TAX INVOICE": (500, 300),
    "Invoice No: 0042": (100, 420),
    "Date: 16/10/2026": (850, 420),
    "Items supplied": (110, 640),
    "Description": (110, 700),
    "Qty": (610, 700),
    "Amount": (860, 700),
    "Paper A4": (110, 760),
    "10": (610, 760),
    "1,250.00": (860, 760),
    "Toner": (110, 820),
    "2": (610, 820),
    "48,000.00": (860, 820),
    "Amount due 49,250.00": (700, 950),
}


class FakeEngine:
    """Returns the lines drawn on the synthetic page, as a real engine would."""

    name = "fake"

    def read(self, img):
        lines = []
        for text, (x, y) in _TEXT.items():
            (tw, th), _ = cv2.getTextSize(text, cv2.FONT_HERSHEY_SIMPLEX, 0.9, 2)
            lines.append(TextLine(text, BBox(x, y - th - 4, x + tw, y + 6)))
        return lines


@pytest.fixture
def invoice_png() -> bytes:
    ok, buf = cv2.imencode(".png", _invoice_image())
    assert ok
    return buf.tobytes()


def test_detect_table_with_merged_header():
    gray = cv2.cvtColor(_invoice_image(), cv2.COLOR_BGR2GRAY)
    tables = detect_tables(gray)
    assert len(tables) == 1
    t = tables[0]
    assert (t.n_rows, t.n_cols) == (4, 3)
    header = [c for c in t.cells if c.row == 0]
    assert len(header) == 1 and header[0].col_span == 3


def test_classify():
    assert classify("TAX INVOICE  Invoice No: 1 Amount due") == "invoice"
    assert classify("Payslip  gross pay  net pay") == "payroll"
    assert classify("hello world") == "general"
    assert recommended_format("payroll") == "xlsx"
    assert recommended_format("invoice") == "docx"


@pytest.mark.parametrize(
    "text,expected",
    [("1,250.00", 1250.0), ("(300)", -300), ("MK 5,000", 5000), ("15%", 0.15), ("Toner", None), ("0042", 42)],
)
def test_parse_number(text, expected):
    assert parse_number(text) == expected


def test_convert_docx(invoice_png):
    result = convert(invoice_png, "invoice.png", "auto", FakeEngine())
    assert result.output_format == "docx"
    assert result.layout.doc_type == "invoice"
    page = result.layout.pages[0]
    kinds = {img.kind for img in page.images}
    assert ImageKind.LOGO in kinds and ImageKind.STAMP in kinds

    doc = Document(io.BytesIO(result.data))
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "TAX INVOICE" in body
    assert "Invoice No: 0042\tDate: 16/10/2026" in body  # same row kept together
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert table.cell(2, 2).text == "1,250.00"
    assert table.cell(0, 0).text == "Items supplied"
    assert len(doc.inline_shapes) == len(page.images)


def test_convert_xlsx_numbers(invoice_png):
    result = convert(invoice_png, "invoice.png", "xlsx", FakeEngine())
    ws = load_workbook(io.BytesIO(result.data))["Page 1"]
    values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert 1250.0 in values and 48000.0 in values
    assert any(str(r).startswith("A") for r in ws.merged_cells.ranges)


def test_api_roundtrip(invoice_png, tmp_path, monkeypatch):
    from app import main

    monkeypatch.setattr(main.settings, "storage_dir", tmp_path)
    main.app.dependency_overrides[main.ocr_engine] = FakeEngine
    try:
        client = TestClient(main.app)
        r = client.post(
            "/v1/conversions",
            files={"file": ("invoice.png", invoice_png, "image/png")},
            data={"output_format": "docx"},
        )
        assert r.status_code == 202, r.text
        job_id = r.json()["id"]
        for _ in range(50):
            meta = client.get(f"/v1/conversions/{job_id}").json()
            if meta["status"] in ("done", "failed"):
                break
            time.sleep(0.1)
        assert meta["status"] == "done", meta
        assert meta["doc_type"] == "invoice" and meta["tables"] == 1
        dl = client.get(f"/v1/conversions/{job_id}/download")
        assert dl.status_code == 200
        Document(io.BytesIO(dl.content))

        assert client.post("/v1/conversions", files={"file": ("x.exe", b"MZ", "application/octet-stream")}).status_code == 415
        assert client.get("/v1/conversions/not-a-uuid").status_code == 404
    finally:
        main.app.dependency_overrides.clear()
